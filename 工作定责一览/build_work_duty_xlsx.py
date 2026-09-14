# -*- coding: utf-8 -*-
"""生成《工作定责一览》Excel（对齐系统管理页岗位口径，未选项目）。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path(__file__).with_name("工作定责一览.xlsx")

# 与 frontend/src/constants/workDuty.ts 及页面未选项目展示一致
# 填写/提交/审核/查看/编辑：岗位文案 +「（待指定到人）」；无此动作为「—」
# 您可办理：系统管理员（超级管理员）口径
ROWS = [
    (
        "立项",
        "项目申报信息与材料",
        "项目联系人、项目负责人（待指定到人）",
        "项目联系人（待指定到人）",
        "项目负责人 → 项目承担部门负责人 → 二级总师 → 单位财务 → 单位科技部长 → 一级总师 → 总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目联系人、项目负责人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/initiation/declaration",
        "草稿可暂存；材料齐套后提交审签；按当前审批节点指定到人；二级总师前须经项目承担部门负责人审核；草稿/驳回可改",
    ),
    (
        "立项",
        "立项支撑材料备案",
        "项目负责人、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位科技部长、总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/initiation/filing",
        "上传本渠道必传原件；齐套后提交归口管理部门；审核通过后台账转实施中；退回补正时可改",
    ),
    (
        "实施",
        "项目基本信息 / 工作内容补录",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位科技部长（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/basic",
        "补齐目标、周期、专业、牵头/参研工作内容、团队；保存并提交二级单位审核；通过后总部科研项目处备案；审核通过前可改，通过后走项目变更",
    ),
    (
        "实施",
        "编制里程碑节点",
        "技术负责人、项目负责人、项目联系人（待指定到人）",
        "技术负责人、项目负责人（待指定到人）",
        "单位科技部长（待指定到人）",
        "项目相关人员（待指定到人）",
        "技术负责人、项目负责人、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/milestone-close?mode=compile",
        "清单齐备后提交存档；二级单位科技部门审核存档",
    ),
    (
        "实施",
        "里程碑销项 / 上传佐证",
        "技术负责人、项目负责人、项目联系人（待指定到人）",
        "项目负责人、技术负责人（待指定到人）",
        "单位科技部长（待指定到人）",
        "项目相关人员（待指定到人）",
        "技术负责人、项目负责人、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/milestone-close",
        "",
    ),
    (
        "实施",
        "计划填报与办结",
        "项目主管、项目负责人（待指定到人）",
        "项目主管、项目负责人（待指定到人）",
        "单位科技部长（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目主管、项目负责人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/plan",
        "办结申请；二级单位管理团队终审，无总部审批",
    ),
    (
        "实施",
        "经费预算填报",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位财务部长 → 总部财务主管（待指定到人）",
        "项目团队、两级财务、管理团队（待指定到人）",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/fund?mode=budget",
        "团队成员可暂存；负责人提交审签；单位财务审核后总部复核备案；草稿/退回可改",
    ),
    (
        "实施",
        "经费核销",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位财务部长（待指定到人）",
        "项目团队、两级财务（待指定到人）",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/fund?mode=writeoff",
        "里程碑闭环后填明细；仅项目负责人提交核销材料；本级核销后同步总部看板；暂存可改",
    ),
    (
        "实施",
        "评估检查",
        "项目负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位科技部长、总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/evaluation",
        "单位上传结论，总部归档",
    ),
    (
        "实施",
        "项目 / 数据变更",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人、项目主管、项目联系人（待指定到人）",
        "单位主管部门 → 总部（重大变更含法务）（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/implement/change",
        "草稿/驳回可改",
    ),
    (
        "验收",
        "交付物维护 / 交付",
        "技术负责人、项目负责人（待指定到人）",
        "项目负责人、技术负责人（待指定到人）",
        "—",
        "项目相关人员（待指定到人）",
        "技术负责人、项目负责人（待指定到人）",
        "您可：填写、提交、编辑",
        "/acceptance/deliverable",
        "本环节无审核动作",
    ),
    (
        "验收",
        "协作单位评价",
        "项目负责人（待指定到人）",
        "项目负责人（待指定到人）",
        "—",
        "项目相关人员（待指定到人）",
        "项目负责人（待指定到人）",
        "您可：填写、提交、编辑",
        "/acceptance/partner",
        "本环节无审核动作",
    ),
    (
        "验收",
        "项目验收申请与办结",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位科技部长、总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、技术负责人、项目主管、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/acceptance/accept",
        "上传本层级验收材料；前置校验通过后提交",
    ),
    (
        "转化",
        "成果转化",
        "项目负责人、项目联系人（待指定到人）",
        "项目负责人（待指定到人）",
        "单位科技部长、总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、项目联系人（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/transform",
        "",
    ),
    (
        "后评价",
        "后评价",
        "项目负责人、单位科技部长、总部处室主管（待指定到人）",
        "项目负责人、单位科技部长（待指定到人）",
        "总部处室（待指定到人）",
        "项目相关人员（待指定到人）",
        "项目负责人、单位科技部长、总部处室主管（待指定到人）",
        "您可：填写、提交、审核、编辑",
        "/post-eval",
        "",
    ),
    (
        "总览",
        "项目台账维护",
        "—",
        "—",
        "—",
        "按数据范围：全公司 / 本单位 / 本人项目（待指定到人）",
        "总部处室、单位科技管理（待指定到人）",
        "您可：编辑",
        "/overview/ledger",
        "本环节无填写/提交/审核；公司领导只读；项目团队改基本信息请走实施填报",
    ),
]

HEADERS = ["阶段", "工作内容", "填写", "提交", "审核", "查看", "编辑", "您可办理"]
EXTRA_HEADERS = HEADERS + ["页面路径", "办理说明"]

NAVY = "0048A0"
BRAND = "0064EF"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FILL = PatternFill("solid", fgColor="E8F1FF")
STAGE_FILL = {
    "立项": PatternFill("solid", fgColor="E6F4FF"),
    "实施": PatternFill("solid", fgColor="F6FFED"),
    "验收": PatternFill("solid", fgColor="FFF7E6"),
    "转化": PatternFill("solid", fgColor="F9F0FF"),
    "后评价": PatternFill("solid", fgColor="FFF1F0"),
    "总览": PatternFill("solid", fgColor="F5F5F5"),
}
MINE_FILL = PatternFill("solid", fgColor="F6FAFF")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color=NAVY)
DESC_FONT = Font(name="微软雅黑", size=10, color="595959")
CELL_FONT = Font(name="微软雅黑", size=10, color="1F1F1F")
MINE_FONT = Font(name="微软雅黑", size=10, bold=True, color=BRAND)
STAGE_FONT = Font(name="微软雅黑", size=10, bold=True, color="1F1F1F")


def style_header_row(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = THIN


def write_data_row(ws: Worksheet, row: int, values: list[str], ncols: int) -> None:
    stage = values[0]
    fill = STAGE_FILL.get(stage, PatternFill("solid", fgColor="FFFFFF"))
    for col, val in enumerate(values[:ncols], 1):
        cell = ws.cell(row, col, val)
        cell.border = THIN
        cell.alignment = WRAP if col != 1 else CENTER
        if col == 1:
            cell.fill = fill
            cell.font = STAGE_FONT
        elif col == ncols and ncols == 8:
            cell.fill = MINE_FILL
            cell.font = MINE_FONT
        else:
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            cell.font = CELL_FONT
    ws.row_dimensions[row].height = 48


def set_widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main() -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "工作定责一览"
    ws.merge_cells("A1:H1")
    ws.merge_cells("A2:H2")
    ws["A1"] = "科研项目信息化管理平台 · 工作定责一览"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
    ws["A2"] = (
        "口径：未选对照项目，按下表岗位展示。进入具体项目后解析为团队实名。"
        "「您可办理」按系统管理员（超级管理员）账号。"
        "公司领导只读；已点名岗位仅该办理人可写；未点名时暂按任职身份办理。"
    )
    ws["A2"].font = DESC_FONT
    ws["A2"].alignment = WRAP
    ws["A2"].fill = TITLE_FILL
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 36
    for col in range(1, 9):
        ws.cell(1, col).fill = TITLE_FILL
        ws.cell(2, col).fill = TITLE_FILL

    for col, h in enumerate(HEADERS, 1):
        ws.cell(3, col, h)
    style_header_row(ws, 3, 8)

    for i, row in enumerate(ROWS):
        write_data_row(ws, 4 + i, list(row), 8)

    set_widths(ws, [10, 24, 28, 24, 42, 28, 28, 26])
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = "A3:H20"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"
    ws.oddHeader.left.text = "COMAC RPM"
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"

    ws2 = wb.create_sheet("含路径与说明")
    ws2.merge_cells("A1:J1")
    ws2["A1"] = "工作定责一览（含页面路径、办理说明）"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].fill = TITLE_FILL
    ws2.row_dimensions[1].height = 28
    for col, h in enumerate(EXTRA_HEADERS, 1):
        ws2.cell(2, col, h)
    style_header_row(ws2, 2, 10)
    for i, row in enumerate(ROWS):
        write_data_row(ws2, 3 + i, list(row), 10)
        ws2.cell(3 + i, 10).alignment = WRAP
    set_widths(ws2, [10, 24, 26, 22, 40, 26, 26, 24, 36, 48])
    ws2.freeze_panes = "C3"
    ws2.auto_filter.ref = "A2:J19"
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.fitToPage = True
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 1

    ws3 = wb.create_sheet("口径说明")
    notes = [
        ["项", "说明"],
        ["页面", "系统管理 → 工作定责一览（/#/system/work-duty）"],
        ["对照项目", "未选项目时按下表岗位口径；选择项目后解析为团队实名办理人"],
        ["填写", "谁可起草/暂存该工作内容"],
        ["提交", "谁可正式提交审签或归档"],
        ["审核", "谁可审核；「—」表示本环节无审核动作"],
        ["查看", "能进入项目即可查看；公司领导只读"],
        ["编辑", "谁可在草稿/驳回/退回补正时修改"],
        ["您可办理", "本表按系统管理员（超级管理员）导出；普通账号随任职身份与是否点名变化"],
        ["待指定到人", "项目团队尚未点名该岗位时的展示；点名后显示「岗位 + 姓名（工号）」"],
        ["数据来源", "给会冉/frontend/src/constants/workDuty.ts"],
        ["导出日期", "2026-09-10"],
    ]
    for r, pair in enumerate(notes, 1):
        ws3.cell(r, 1, pair[0])
        ws3.cell(r, 2, pair[1])
        ws3.cell(r, 1).border = THIN
        ws3.cell(r, 2).border = THIN
        ws3.cell(r, 1).font = STAGE_FONT if r > 1 else HEAD_FONT
        ws3.cell(r, 2).font = CELL_FONT if r > 1 else HEAD_FONT
        ws3.cell(r, 1).fill = HEAD_FILL if r == 1 else TITLE_FILL
        ws3.cell(r, 2).fill = HEAD_FILL if r == 1 else PatternFill("solid", fgColor="FFFFFF")
        ws3.cell(r, 1).alignment = CENTER
        ws3.cell(r, 2).alignment = WRAP
        if r == 1:
            ws3.cell(r, 1).font = HEAD_FONT
            ws3.cell(r, 2).font = HEAD_FONT
        ws3.row_dimensions[r].height = 22
    set_widths(ws3, [16, 88])
    ws3.row_dimensions[1].height = 24

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
